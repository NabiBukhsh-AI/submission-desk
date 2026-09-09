"""Word documents: paragraphs, and tables flattened into rows.

A CV in a table is common and reads badly if the cells are concatenated without
a separator: "Senior EngineerAcme2021" is not a sentence anyone can quote. Cells
are joined with " | ", and that choice is recorded in the normalisation profile
id, because a quotation validated against one delimiter will not match text
assembled with another.

A DOCX has no pages until it is rendered, so the whole document is one page
here. That is honest rather than convenient: a page number invented at this
stage would send a reviewer to the wrong place in their own file.
"""

from __future__ import annotations

import io

import docx

from domain.contracts.enums import ExtractionMethod
from domain.ports.extraction import ExtractionFailed, PageText

#: Between table cells. Part of the profile id, because span validation depends
#: on the text being assembled the same way every time.
CELL_DELIMITER = " | "


def extract_pages(data: bytes) -> list[PageText]:
    """Read a Word document into a single page of ordered blocks."""
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:
        raise ExtractionFailed(
            "This Word document could not be opened. Ask the candidate to send it again.",
            error_code="CORRUPT_FILE",
        ) from error

    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(CELL_DELIMITER.join(cells))

    text = "\n".join(parts)
    blocks = tuple(
        (0.0, float(index), 1.0, float(index) + 1.0, part) for index, part in enumerate(parts)
    )

    return [
        PageText(
            page_number=1,
            text=text,
            method=ExtractionMethod.DOCX,
            blocks=blocks,
            width=1.0,
            height=float(len(parts) or 1),
        )
    ]
